"""Вывод результата на экран и в Excel."""

from decimal import Decimal

import stavki


KOLONKI = [
    ("Наименование", "naimenovanie"),
    ("Страна", "strana"),
    ("Кол-во", "kolichestvo"),
    ("Цена за ед. (валюта)", "tsena"),
    ("Вид", "vid_nazv"),
    ("Стоимость (цена поставщика), сом", "stoimost"),
    ("База НДС/НДН, сом", "tam_stoimost"),
    ("Пошлина %", "poshlina_stavka"),
    ("Пошлина, сом", "poshlina"),
    ("Акциз, сом", "akciz"),
    ("НДН, сом", "ndn"),
    ("НДС импорт / за нерезидента, сом", "nds"),
    ("Тамож. сбор, сом", "sbor"),
    ("Прочие, сом", "prochie"),
    ("Себестоимость, сом", "sebestoimost"),
    ("Себест. за ед., сом", "sebest_ed"),
    ("Цена продажи за ед., сом", "tsena_prodazhi_ed"),
    ("Выручка без НДС, сом", "vyruchka"),
    ("Налог с выручки, сом", "nalog_vyruchka"),
    ("Прибыль, сом", "pribyl"),
]

ITOGO_KLYUCHI = ["stoimost", "tam_stoimost", "poshlina", "akciz", "ndn", "nds", "sbor",
                 "prochie", "sebestoimost", "vyruchka", "nalog_vyruchka", "pribyl"]


def pechat(tovary, platelshik_nds):
    print("\n" + "=" * 60)
    print("РЕЗУЛЬТАТ")
    print("=" * 60)
    for t in tovary:
        if t.get("vid") == "usluga":
            print(f"\nУСЛУГА {t['naimenovanie']} ({t['strana']}), {t['kolichestvo']} ед.")
            print(f"  {'Цена поставщика:':<24}{t['stoimost']:>14,.2f} {stavki.VAL_NAZV}")
            print(f"  {'База (НДН ' + ('сверху)' if t['ndn_sverhu'] else 'удерж.)') + ':':<24}"
                  f"{t['tam_stoimost']:>14,.2f} {stavki.VAL_NAZV}")
            print(f"  {'НДН (' + str(t['ndn_stavka']) + '%):':<24}{t['ndn']:>14,.2f} {stavki.VAL_NAZV}")
            print(f"  {'НДС нерезидент (' + str(stavki.NDS_STAVKA) + '%):':<24}{t['nds']:>14,.2f} {stavki.VAL_NAZV}"
                  + ("  (к зачёту)" if platelshik_nds else "  (в расход)"))
            print(f"  СЕБЕСТОИМОСТЬ:          {t['sebestoimost']:>14,.2f} {stavki.VAL_NAZV}")
            print(f"  цена продажи за ед.:    {t['tsena_prodazhi_ed']:>14,.2f} {stavki.VAL_NAZV}")
            print(f"  ПРИБЫЛЬ:                {t['pribyl']:>14,.2f} {stavki.VAL_NAZV}")
            continue
        print(f"\n{t['naimenovanie']} ({t['strana']}), {t['kolichestvo']} шт.")
        print(f"  Таможенная стоимость:   {t['tam_stoimost']:>14,.2f} {stavki.VAL_NAZV}")
        print(f"  Пошлина ({t['poshlina_stavka']}%):        {t['poshlina']:>14,.2f} {stavki.VAL_NAZV}")
        print(f"  Акциз:                  {t['akciz']:>14,.2f} {stavki.VAL_NAZV}")
        print(f"  НДС на импорт ({stavki.NDS_STAVKA}%):   {t['nds']:>14,.2f} {stavki.VAL_NAZV}"
              + ("  (к зачёту)" if platelshik_nds else "  (в себестоимость)"))
        print(f"  Тамож. сбор + прочие:   {t['sbor'] + t['prochie']:>14,.2f} {stavki.VAL_NAZV}")
        print(f"  СЕБЕСТОИМОСТЬ:          {t['sebestoimost']:>14,.2f} {stavki.VAL_NAZV}")
        print(f"  за единицу:             {t['sebest_ed']:>14,.2f} {stavki.VAL_NAZV}")
        print(f"  цена продажи за ед.:    {t['tsena_prodazhi_ed']:>14,.2f} {stavki.VAL_NAZV}"
              + ("  (с НДС)" if platelshik_nds else ""))
        print(f"  выручка без НДС:        {t['vyruchka']:>14,.2f} {stavki.VAL_NAZV}")
        print(f"  налог с выручки:        {t['nalog_vyruchka']:>14,.2f} {stavki.VAL_NAZV}")
        print(f"  ПРИБЫЛЬ:                {t['pribyl']:>14,.2f} {stavki.VAL_NAZV}")

    print("\n" + "-" * 60)
    for nazv, k in [("Стоимость (цена)", "stoimost"), ("Пошлина", "poshlina"),
                    ("НДН", "ndn"),
                    ("Акциз", "akciz"), ("НДС импорт/нерезидент", "nds"),
                    ("Себестоимость партии", "sebestoimost"),
                    ("Выручка без НДС", "vyruchka"), ("Налог с выручки", "nalog_vyruchka"),
                    ("ПРИБЫЛЬ", "pribyl")]:
        print(f"ИТОГО {nazv + ':':<24}{sum(t[k] for t in tovary):>14,.2f} {stavki.VAL_NAZV}")
    print(f"ИТОГО в бюджет (пошлина+акциз+НДС+сбор+НДН): "
          f"{sum(t['k_oplate'] for t in tovary):,.2f} {stavki.VAL_NAZV}")


def sohranit_excel(tovary, shapka, imya_fayla):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("\nopenpyxl не установлен — Excel не сохранён. "
              "Установите: .\\.venv\\Scripts\\python.exe -m pip install openpyxl")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Расчёт"
    for stroka in shapka:
        ws.append(stroka)
    ws.append([])
    ws.append([nazv.replace("сом", stavki.VAL_NAZV) for nazv, _ in KOLONKI])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for t in tovary:
        t["vid_nazv"] = "услуга" if t.get("vid") == "usluga" else "товар"
        ws.append([float(t[k]) if isinstance(t[k], Decimal) else t[k] for _, k in KOLONKI])
    itogo = ["ИТОГО"] + [""] * (len(KOLONKI) - 1)
    for i, (_, k) in enumerate(KOLONKI):
        if k in ITOGO_KLYUCHI:
            itogo[i] = float(sum(t[k] for t in tovary))
    ws.append(itogo)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for kol in ws.columns:
        ws.column_dimensions[kol[0].column_letter].width = 18
    ws.column_dimensions["A"].width = 30
    wb.save(imya_fayla)
    print(f"\nСохранено в Excel: {imya_fayla}")
